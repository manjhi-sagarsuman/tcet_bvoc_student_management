from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from .models import Student

HEADERS = ["Student ID", "Full Name", "B.Voc Branch", "Year", "Semester",
           "Academic Year", "Email", "Phone", "Status"]

class ExcelStudentRepository:
    def __init__(self, path="data/students.xlsx"):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            self._create()

    def _create(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for col, width in zip("ABCDEFGHI", [18, 26, 42, 14, 14, 18, 30, 18, 16]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:I1"
        wb.save(self.path)
        wb.close()

    def all(self):
        wb = load_workbook(self.path, data_only=True)
        ws = wb["Students"]
        records = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if not values or not values[0]:
                continue
            values = list(values) + [""] * (9 - len(values))
            records.append(Student(*[str(v or "") for v in values]))
        wb.close()
        return records

    def save(self, student):
        wb = load_workbook(self.path)
        ws = wb["Students"]
        target_row = None
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, 1).value).strip() == student.student_id:
                target_row = row
                break
        values = [student.student_id, student.full_name, student.branch, student.year,
                  student.semester, student.academic_year, student.email, student.phone, student.status]
        if target_row:
            for col, value in enumerate(values, 1):
                ws.cell(target_row, col, value)
        else:
            ws.append(values)
        wb.save(self.path)
        wb.close()
        return "updated" if target_row else "added"

    def delete(self, student_id):
        wb = load_workbook(self.path)
        ws = wb["Students"]
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row, 1).value).strip() == student_id:
                ws.delete_rows(row, 1)
                wb.save(self.path)
                wb.close()
                return True
        wb.save(self.path)
        wb.close()
        return False
